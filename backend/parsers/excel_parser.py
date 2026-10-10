"""Excel workbook parsing for invoice and bank-statement spreadsheets."""

import logging
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any

from backend.parsers.bank_csv_parser import parse_bank_csv_text
from backend.schemas.evidence import Evidence, EvidenceProvenance
from backend.services.financial_field_extractor import _normalize_number

LOGGER = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
DEFAULT_LANGUAGE = "eng"

_INVOICE_MARKERS = {
    "invoice_id", "invoice_number", "seller_gstin", "buyer_gstin", "invoice_date",
    "seller_name", "buyer_name",
}
_BANK_MARKERS = {"transaction_date", "txn_date", "narration", "debit", "credit", "running_balance"}
_GENERIC_MARKERS = {"file_name", "invoice_category", "data_type"}


def validate_excel_path(file_path: str | os.PathLike[str]) -> Path:
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Excel file does not exist: {path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported Excel format '{path.suffix}'. Supported formats: {supported}")
    return path


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _header_key(value: Any) -> str:
    text = _cell_text(value).casefold()
    return "".join(character if character.isalnum() else "_" for character in text).strip("_")


def _load_rows(path: Path) -> tuple[str, list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires openpyxl. Install dependencies with: pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        best_rows: list[list[str]] = []
        best_sheet = ""
        for worksheet in workbook.worksheets:
            rows = [[_cell_text(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(row)]
            if len(rows) > len(best_rows):
                best_rows = rows
                best_sheet = worksheet.title
        if not best_rows:
            raise ValueError(f"Excel workbook '{path.name}' contains no data")
        return best_sheet, best_rows
    finally:
        workbook.close()


def _find_header(rows: list[list[str]]) -> tuple[int, list[str]]:
    best_index = 0
    best_headers: list[str] = []
    best_score = -1
    for index, row in enumerate(rows[:30]):
        headers = [_header_key(value) for value in row]
        score = sum(
            1 for header in headers
            if header in _INVOICE_MARKERS or header in _BANK_MARKERS
            or header in _GENERIC_MARKERS
            or header in {"date", "description", "amount", "balance", "seller_name", "buyer_name"}
        )
        if score > best_score:
            best_index, best_headers, best_score = index, headers, score
    if best_score <= 0:
        raise ValueError("Could not identify an Excel header row")
    return best_index, best_headers


def _column(headers: list[str], *names: str) -> int | None:
    for name in names:
        key = _header_key(name)
        if key in headers:
            return headers.index(key)
    return None


def _value(row: list[str], index: int | None) -> str | None:
    if index is None or index >= len(row) or not row[index]:
        return None
    return row[index]


def _excel_amount(value: str) -> str:
    normalized = _normalize_number(value)
    if "." not in normalized:
        return f"{normalized}.00"
    whole, fraction = normalized.split(".", 1)
    return f"{whole}.{fraction[:2].ljust(2, '0')}"


def _parse_invoice_workbook(
    path: Path,
    sheet_name: str,
    headers: list[str],
    rows: list[list[str]],
    header_index: int,
    language: str,
    business_id: str | None,
    persist_to_database: bool,
    file_type: str,
) -> Evidence:
    column_names = {
        "invoice_number": _column(headers, "invoice_id", "invoice_number", "invoice_no"),
        "date": _column(headers, "invoice_date", "date"),
        "due_date": _column(headers, "due_date", "payment_due_date"),
        "invoice_type": _column(headers, "invoice_type", "document_type"),
        "seller_name": _column(headers, "seller_name", "vendor", "supplier"),
        "seller_gstin": _column(headers, "seller_gstin", "seller_gst_in", "gstin"),
        "buyer_name": _column(headers, "buyer_name", "customer", "client"),
        "buyer_gstin": _column(headers, "buyer_gstin", "buyer_gst_in"),
        "place_of_supply": _column(headers, "place_of_supply", "state"),
        "description": _column(headers, "item_description", "description", "service_description", "file_name", "invoice_category"),
        "quantity": _column(headers, "quantity", "qty"),
        "unit_price": _column(headers, "unit_price", "price"),
        "subtotal": _column(headers, "taxable_value", "subtotal", "taxable_amount"),
        "cgst_total": _column(headers, "cgst", "cgst_amount", "cgst_total"),
        "sgst_total": _column(headers, "sgst", "sgst_amount", "sgst_total"),
        "igst_total": _column(headers, "igst", "igst_amount", "igst_total"),
        "invoice_total": _column(headers, "grand_total", "invoice_total", "total_amount", "total"),
        "currency": _column(headers, "currency"),
    }
    data_rows = [row for row in rows[header_index + 1:] if any(row)]
    if not data_rows:
        raise ValueError(f"Excel sheet '{sheet_name}' contains no records")

    first = data_rows[0]
    fields: dict[str, Any] = {
        key: _value(first, index)
        for key, index in column_names.items()
        if _value(first, index) is not None
    }
    for key in ("subtotal", "cgst_total", "sgst_total", "igst_total", "invoice_total", "unit_price"):
        if fields.get(key):
            fields[key] = _excel_amount(fields[key])
    fields["amount"] = fields.get("invoice_total")
    fields["party_name"] = fields.get("seller_name")
    fields["invoice_date"] = fields.get("date")
    fields["items"] = []
    for item_number, row in enumerate(data_rows, start=1):
        item = {
            "line_number": item_number,
            "description": _value(row, column_names["description"]),
            "quantity": _value(row, column_names["quantity"]),
            "unit_price": _value(row, column_names["unit_price"]),
            "line_total": _value(row, column_names["invoice_total"]),
        }
        if any(value is not None for value in item.values() if value != item_number):
            fields["items"].append(item)

    lines = [" | ".join(f"{headers[index]}: {value}" for index, value in enumerate(row) if value) for row in data_rows]
    extracted_text = "\n".join(lines)
    provenance = []
    for field_name in ("invoice_number", "date", "seller_name", "seller_gstin", "buyer_name", "buyer_gstin", "invoice_total"):
        value = fields.get(field_name)
        if value:
            provenance.append(EvidenceProvenance(
                field_name=field_name,
                value=str(value),
                source_text=lines[0],
                line_number=header_index + 2,
                page_number=1,
                extraction_method="excel_column",
            ))
    evidence = Evidence(
        source_type="excel",
        file_name=path.name,
        extracted_text=extracted_text,
        language=language,
        processing_status="success",
        page_count=1,
        provenance=provenance,
        **fields,
    )
    if persist_to_database:
        from backend.database.supabase_repository import persist_evidence
        ids = persist_evidence(path, evidence, business_id, file_type)
        evidence = evidence.model_copy(update={
            "database_file_id": ids["file_id"],
            "database_evidence_id": ids["evidence_id"],
            "database_invoice_id": ids.get("invoice_id"),
        })
    return evidence


def parse_excel_file(
    file_path: str | os.PathLike[str],
    language: str = DEFAULT_LANGUAGE,
    business_id: str | None = None,
    persist_to_database: bool = False,
    file_type: str = "invoice",
) -> Evidence:
    """Parse an Excel workbook as an invoice table or bank statement."""
    path = Path(file_path)
    try:
        validate_excel_path(path)
        sheet_name, rows = _load_rows(path)
        header_index, headers = _find_header(rows)
        markers = set(headers)
        if markers & _BANK_MARKERS and not markers & _INVOICE_MARKERS:
            csv_text = "\n".join(",".join(row) for row in rows[header_index:])
            return Evidence.model_validate(
                parse_bank_csv_text(
                    csv_text,
                    source_name=path.name,
                    language=language,
                    business_id=business_id,
                    persist_to_database=persist_to_database,
                    file_type="bank_statement",
                ).model_dump()
            )
        return _parse_invoice_workbook(
            path, sheet_name, headers, rows, header_index, language,
            business_id, persist_to_database, file_type,
        )
    except Exception as exc:
        LOGGER.exception("Excel parsing failed for %s", path)
        return Evidence(
            source_type="excel",
            file_name=path.name,
            language=language,
            processing_status="failed",
            error=str(exc),
        )
